#!/usr/bin/env python3
"""One 16:9 slide: app render loop to display, Worker steps highlighted."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# Widescreen 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK = RGBColor(0x1C, 0x1B, 0x16)
MUTED = RGBColor(0x5C, 0x58, 0x4E)
LINE = RGBColor(0xC8, 0xC1, 0xB0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OTHER_FILL = RGBColor(0xF7, 0xF4, 0xEE)
WORKER_FILL = RGBColor(0xD6, 0xE8, 0xF8)
WORKER_LINE = RGBColor(0x15, 0x65, 0xC0)
SF_FILL = RGBColor(0xF0, 0xE4, 0xF6)
SF_LINE = RGBColor(0x6A, 0x1B, 0x9A)
TITLE_FILL = RGBColor(0x1C, 0x1B, 0x16)
SCREEN_FILL = RGBColor(0xE4, 0xF0, 0xE2)
SCREEN_LINE = RGBColor(0x2E, 0x7D, 0x32)

FONT = "Microsoft YaHei"


def set_run_font(run, size_pt, bold=False, color=INK, font=FONT):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font)
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = rPr.makeelement(qn("a:latin"), {})
        rPr.append(latin)
    latin.set("typeface", font)


def add_textbox(slide, l, t, w, h, text, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size, bold, color)
    return box


def add_box(slide, l, t, w, h, who, title, fill, line, who_color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.25)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(70000)
    tf.margin_right = Emu(70000)
    tf.margin_top = Emu(40000)
    tf.margin_bottom = Emu(30000)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT

    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = who
    set_run_font(r0, 9, True, who_color)

    p1 = tf.add_paragraph()
    p1.alignment = PP_ALIGN.LEFT
    r1 = p1.add_run()
    r1.text = title
    set_run_font(r1, 12, True, INK)
    return shape


def add_arrow_right(slide, x, y, w=Inches(0.22), h=Inches(0.14)):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0x9A, 0x94, 0x86)
    sh.line.fill.background()
    return sh


def add_arrow_down(slide, x, y, w=Inches(0.16), h=Inches(0.20)):
    sh = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0x9A, 0x94, 0x86)
    sh.line.fill.background()
    return sh


def add_legend_chip(slide, l, t, fill, line, label):
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(1.55), Inches(0.28))
    chip.adjustments[0] = 0.3
    chip.fill.solid()
    chip.fill.fore_color.rgb = fill
    chip.line.color.rgb = line
    tf = chip.text_frame
    tf.word_wrap = False
    tf.margin_left = Emu(50000)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    set_run_font(r, 10, True, INK)


def place_row(slide, boxes, y, left, right, box_h, who_color_for):
    n = len(boxes)
    gap = Inches(0.28)
    usable = right - left - gap * (n - 1)
    bw = usable / n
    xs = []
    for i, spec in enumerate(boxes):
        x = left + i * (bw + gap)
        fill, line, who_c = who_color_for(spec["kind"])
        add_box(slide, x, y, bw, box_h, spec["who"], spec["title"], fill, line, who_c)
        xs.append((x, bw))
        if i < n - 1:
            add_arrow_right(slide, x + bw + Inches(0.03), y + box_h / 2 - Inches(0.07))
    return xs


def kind_style(kind):
    if kind == "app_worker":
        return WORKER_FILL, WORKER_LINE, WORKER_LINE
    if kind == "sf_worker":
        return SF_FILL, SF_LINE, SF_LINE
    if kind == "sf":
        return SF_FILL, SF_LINE, SF_LINE
    if kind == "screen":
        return SCREEN_FILL, SCREEN_LINE, SCREEN_LINE
    return OTHER_FILL, LINE, MUTED


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0xFB, 0xF9, 0xF4)
    bg.line.fill.background()

    add_textbox(
        slide,
        Inches(0.35),
        Inches(0.18),
        Inches(9.2),
        Inches(0.42),
        "从应用渲染循环到上屏：Worker 在哪",
        size=24,
        bold=True,
    )
    add_textbox(
        slide,
        Inches(0.35),
        Inches(0.55),
        Inches(9.2),
        Inches(0.32),
        "redroid guest · vulkan.pastel.so · 一帧里 Worker 只出现两次（应用池 + SurfaceFlinger 池）",
        size=12,
        color=MUTED,
    )

    add_legend_chip(slide, Inches(8.35), Inches(0.22), WORKER_FILL, WORKER_LINE, "蓝：应用 Worker")
    add_legend_chip(slide, Inches(10.02), Inches(0.22), SF_FILL, SF_LINE, "紫：SF Worker")
    add_legend_chip(slide, Inches(11.68), Inches(0.22), OTHER_FILL, LINE, "灰：其他线程")

    left, right = Inches(0.35), Inches(12.98)
    h = Inches(0.95)

    row1 = [
        {"kind": "other", "who": "显示", "title": "1. Vsync 到了"},
        {"kind": "other", "who": "应用 UI 线程", "title": "2. 改界面、记 DisplayList"},
        {"kind": "other", "who": "RenderThread / HWUI", "title": "3. 变成 Vulkan 命令"},
        {"kind": "other", "who": "Queue 线程", "title": "4. vkQueueSubmit"},
        {"kind": "other", "who": "Renderer", "title": "5. 准备 / JIT Routine"},
    ]
    y1 = Inches(1.00)
    xs1 = place_row(slide, row1, y1, left, right, h, kind_style)
    add_arrow_down(slide, xs1[-1][0] + xs1[-1][1] / 2 - Inches(0.08), y1 + h + Inches(0.02))

    add_textbox(
        slide,
        left,
        Inches(2.18),
        Inches(4.5),
        Inches(0.24),
        "应用进程 · vulkan.pastel.so",
        size=11,
        bold=True,
        color=WORKER_LINE,
    )

    row2 = [
        {"kind": "app_worker", "who": "应用 Worker", "title": "6. batch 入队  marl::schedule"},
        {"kind": "app_worker", "who": "应用 Worker", "title": "7. 顶点 → 图元 → 像素"},
        {"kind": "other", "who": "CPU 内存", "title": "8. 颜色在 vk::Image"},
        {"kind": "other", "who": "Android 导出", "title": "9. memcpy 到 AHB，queueBuffer"},
    ]
    y2 = Inches(2.40)
    xs2 = place_row(slide, row2, y2, left, right, h, kind_style)
    add_arrow_down(slide, xs2[-1][0] + xs2[-1][1] / 2 - Inches(0.08), y2 + h + Inches(0.02))

    add_textbox(
        slide,
        left,
        Inches(3.58),
        Inches(5.5),
        Inches(0.24),
        "SurfaceFlinger · 另一份 ICD，另一套 ThreadCount",
        size=11,
        bold=True,
        color=SF_LINE,
    )

    row3 = [
        {"kind": "sf", "who": "SF 主线程", "title": "10. 排图层、Client 合成"},
        {"kind": "sf", "who": "SF Queue", "title": "11. 自己的 vkQueueSubmit"},
        {"kind": "sf_worker", "who": "SF Worker", "title": "12. 采样各层 + 片元写出"},
        {"kind": "screen", "who": "显示 / redroid", "title": "13. 上屏"},
    ]
    y3 = Inches(3.80)
    place_row(slide, row3, y3, left, right, h, kind_style)

    # Bottom callout: what happens inside a Worker
    note = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.35),
        Inches(5.00),
        Inches(12.63),
        Inches(2.18),
    )
    note.adjustments[0] = 0.04
    note.fill.solid()
    note.fill.fore_color.rgb = WHITE
    note.line.color.rgb = LINE

    add_textbox(
        slide,
        Inches(0.55),
        Inches(5.12),
        Inches(12.2),
        Inches(0.32),
        "一块 batch 在 Worker 里（步骤 7 / 12）",
        size=14,
        bold=True,
    )
    add_textbox(
        slide,
        Inches(0.55),
        Inches(5.42),
        Inches(12.2),
        Inches(0.85),
        "DrawCall::run  →  marl::schedule(batch)  →  processVertices（VertexRoutine）  →  processPrimitives（SetupRoutine）  →  processPixels（PixelRoutine + sampler，四像素一组）\n"
        "没活：spinForWork 空转一会儿，再 wait() 睡觉。挡住了：suspend，切纤程（Fiber），同一条 OS 线程接着干别的。",
        size=13,
        color=INK,
    )
    add_textbox(
        slide,
        Inches(0.55),
        Inches(6.28),
        Inches(12.2),
        Inches(0.72),
        "不是 Worker：UI 线程、RenderThread、Queue<…>、SF 主线程。绑核只限制坐哪几把椅子；条数认 cwd 的 ThreadCount，否则 min(看见的逻辑核, 16)。应用和 SF 各走一遍这条链。",
        size=12,
        color=MUTED,
    )

    out = "/workspace/docs/MarlWorkerFrame.pptx"
    prs.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
